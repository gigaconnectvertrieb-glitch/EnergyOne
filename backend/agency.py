import base64
import csv
import io
import json
import os
import smtplib
import uuid
from datetime import date, datetime, timedelta
from email.mime.text import MIMEText
from typing import Literal, Optional

from fastapi import Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, EmailStr, Field
from sqlalchemy import Boolean, Date, DateTime, Float, ForeignKey, Integer, String, Text, func, or_, select
from sqlalchemy.orm import Mapped, Session, mapped_column

import pyotp

from .app import Base, Activity, Customer, CustomerHistory, CustomerIn, Employee, EmployeeIn, JobApplication, MasterKeyIn, Module, MODULE_SEED, STORAGE, Task, TaskIn, Team, TeamMember, app, current, admin, admin_or_lead, db, log, module_enabled, notify, serialize, serialize_employee, create_customer, create_task, create_employee, reset_totp, rotate_master_key, make_pdf, notify_update, require_module, visible_employee_ids


class ScheduleEntry(Base):
    __tablename__ = "mitarbeiter_planung"
    id: Mapped[int] = mapped_column(primary_key=True)
    employee_id: Mapped[int] = mapped_column(ForeignKey("mitarbeiter.id"))
    starts_at: Mapped[datetime] = mapped_column(DateTime)
    ends_at: Mapped[datetime] = mapped_column(DateTime)
    kind: Mapped[str] = mapped_column(String(30), default="arbeit")
    note: Mapped[str] = mapped_column(Text, default="")


class News(Base):
    __tablename__ = "news"
    id: Mapped[int] = mapped_column(primary_key=True)
    title: Mapped[str] = mapped_column(String(200))
    text: Mapped[str] = mapped_column(Text)
    author_id: Mapped[int] = mapped_column(ForeignKey("mitarbeiter.id"))
    published_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)
    important: Mapped[bool] = mapped_column(Boolean, default=False)


class Incentive(Base):
    __tablename__ = "incentives"
    id: Mapped[int] = mapped_column(primary_key=True)
    name: Mapped[str] = mapped_column(String(150))
    description: Mapped[str] = mapped_column(Text, default="")
    minimum_contracts: Mapped[int] = mapped_column(Integer, default=0)
    reward_eur: Mapped[float] = mapped_column(Float)
    active: Mapped[bool] = mapped_column(Boolean, default=True)



class Expense(Base):
    __tablename__ = "ausgaben"
    id: Mapped[int] = mapped_column(primary_key=True)
    category: Mapped[str] = mapped_column(String(100))
    description: Mapped[str] = mapped_column(Text, default="")
    amount: Mapped[float] = mapped_column(Float)
    spent_on: Mapped[date] = mapped_column(Date, default=date.today)
    created_by: Mapped[int] = mapped_column(ForeignKey("mitarbeiter.id"))


class SalesGoal(Base):
    __tablename__ = "vertriebsziele"
    id: Mapped[int] = mapped_column(primary_key=True)
    employee_id: Mapped[Optional[int]] = mapped_column(ForeignKey("mitarbeiter.id"), nullable=True)
    team_id: Mapped[Optional[int]] = mapped_column(ForeignKey("teams.id"), nullable=True)
    period_start: Mapped[date] = mapped_column(Date)
    period_end: Mapped[date] = mapped_column(Date)
    target_contracts: Mapped[int] = mapped_column(Integer, default=0)
    target_revenue: Mapped[float] = mapped_column(Float, default=0)


class StornoStatus(Base):
    __tablename__ = "storno_status"
    id: Mapped[int] = mapped_column(primary_key=True)
    employee_id: Mapped[int] = mapped_column(ForeignKey("mitarbeiter.id"), unique=True)
    alert: Mapped[bool] = mapped_column(Boolean, default=False)
    rate: Mapped[float] = mapped_column(Float, default=0)
    updated_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)


class Document(Base):
    __tablename__ = "dokumente"
    id: Mapped[int] = mapped_column(primary_key=True)
    owner_employee_id: Mapped[Optional[int]] = mapped_column(ForeignKey("mitarbeiter.id"), nullable=True)
    category: Mapped[str] = mapped_column(String(40))
    filename: Mapped[str] = mapped_column(String(255))
    storage_name: Mapped[str] = mapped_column(String(255))
    amount: Mapped[Optional[float]] = mapped_column(Float, nullable=True)
    paid: Mapped[Optional[bool]] = mapped_column(Boolean, nullable=True)
    expires_on: Mapped[Optional[date]] = mapped_column(Date, nullable=True)
    note: Mapped[str] = mapped_column(Text, default="")
    uploaded_by: Mapped[int] = mapped_column(ForeignKey("mitarbeiter.id"))
    uploaded_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)


class Provider(Base):
    __tablename__ = "anbieter"
    id: Mapped[int] = mapped_column(primary_key=True)
    name: Mapped[str] = mapped_column(String(150), unique=True)
    street: Mapped[Optional[str]] = mapped_column(String(200), nullable=True)
    postal_code: Mapped[Optional[str]] = mapped_column(String(10), nullable=True)
    city: Mapped[Optional[str]] = mapped_column(String(120), nullable=True)
    phone: Mapped[Optional[str]] = mapped_column(String(50), nullable=True)
    email: Mapped[Optional[str]] = mapped_column(String(255), nullable=True)
    contact_person: Mapped[Optional[str]] = mapped_column(String(150), nullable=True)
    notes: Mapped[str] = mapped_column(Text, default="")
    active: Mapped[bool] = mapped_column(Boolean, default=True)
    is_own_product: Mapped[bool] = mapped_column(Boolean, default=False)


class Tariff(Base):
    __tablename__ = "tarife"
    id: Mapped[int] = mapped_column(primary_key=True)
    provider_id: Mapped[int] = mapped_column(ForeignKey("anbieter.id"))
    name: Mapped[str] = mapped_column(String(200))
    external_id: Mapped[Optional[str]] = mapped_column(String(30), nullable=True)
    active: Mapped[bool] = mapped_column(Boolean, default=True)
    product_type: Mapped[Optional[str]] = mapped_column(String(20), nullable=True)
    base_price_monthly: Mapped[Optional[float]] = mapped_column(Float, nullable=True)
    price_per_kwh: Mapped[Optional[float]] = mapped_column(Float, nullable=True)
    contract_term_months: Mapped[Optional[int]] = mapped_column(Integer, nullable=True)
    description: Mapped[Optional[str]] = mapped_column(Text, nullable=True)


class CommissionBracket(Base):
    __tablename__ = "provisions_staffeln"
    id: Mapped[int] = mapped_column(primary_key=True)
    tariff_id: Mapped[int] = mapped_column(ForeignKey("tarife.id"))
    tier: Mapped[int] = mapped_column(Integer)
    usage_from: Mapped[float] = mapped_column(Float, default=0)
    usage_to: Mapped[Optional[float]] = mapped_column(Float, nullable=True)
    commission_amount: Mapped[float] = mapped_column(Float, default=0)
    commission_per_kwh: Mapped[Optional[float]] = mapped_column(Float, nullable=True)
    active: Mapped[bool] = mapped_column(Boolean, default=True)


class OwnProductOrder(Base):
    __tablename__ = "eigene_bestellungen"
    id: Mapped[int] = mapped_column(primary_key=True)
    tariff_id: Mapped[int] = mapped_column(ForeignKey("tarife.id"))
    name: Mapped[str] = mapped_column(String(150))
    email: Mapped[str] = mapped_column(String(255))
    phone: Mapped[Optional[str]] = mapped_column(String(50), nullable=True)
    postal_code: Mapped[str] = mapped_column(String(10))
    usage_kwh: Mapped[float] = mapped_column(Float, default=0)
    estimated_monthly_price: Mapped[Optional[float]] = mapped_column(Float, nullable=True)
    status: Mapped[str] = mapped_column(String(20), default="interessent")
    employee_id: Mapped[Optional[int]] = mapped_column(ForeignKey("mitarbeiter.id"), nullable=True)
    created_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)


class ClosureEntry(Base):
    __tablename__ = "abschluss_meldungen"
    id: Mapped[int] = mapped_column(primary_key=True)
    employee_id: Mapped[int] = mapped_column(ForeignKey("mitarbeiter.id"))
    customer_id: Mapped[Optional[int]] = mapped_column(ForeignKey("kunden.id"), nullable=True)
    customer_name: Mapped[str] = mapped_column(String(200))
    contract_number: Mapped[Optional[str]] = mapped_column(String(50), nullable=True)
    product: Mapped[str] = mapped_column(String(20), default="strom")
    customer_kind: Mapped[str] = mapped_column(String(20), default="privat")
    usage_kwh: Mapped[float] = mapped_column(Float, default=0)
    completed_on: Mapped[date] = mapped_column(Date, default=date.today)
    status: Mapped[str] = mapped_column(String(20), default="eingereicht")
    provider_id: Mapped[Optional[int]] = mapped_column(ForeignKey("anbieter.id"), nullable=True)
    bracket_id: Mapped[Optional[int]] = mapped_column(ForeignKey("provisions_staffeln.id"), nullable=True)
    expected_commission: Mapped[float] = mapped_column(Float, default=0)
    note: Mapped[str] = mapped_column(Text, default="")
    reviewed_by: Mapped[Optional[int]] = mapped_column(ForeignKey("mitarbeiter.id"), nullable=True)
    reviewed_at: Mapped[Optional[datetime]] = mapped_column(DateTime, nullable=True)


class DailyPerformance(Base):
    __tablename__ = "tagesmeldungen"
    id: Mapped[int] = mapped_column(primary_key=True)
    employee_id: Mapped[int] = mapped_column(ForeignKey("mitarbeiter.id"))
    entry_date: Mapped[date] = mapped_column(Date)
    contracts: Mapped[int] = mapped_column(Integer, default=0)
    cancellations: Mapped[int] = mapped_column(Integer, default=0)
    note: Mapped[str] = mapped_column(Text, default="")
    created_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)
    updated_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)


class Training(Base):
    __tablename__ = "schulungen"
    id: Mapped[int] = mapped_column(primary_key=True)
    title: Mapped[str] = mapped_column(String(200))
    description: Mapped[str] = mapped_column(Text, default="")
    starts_at: Mapped[datetime] = mapped_column(DateTime)
    ends_at: Mapped[datetime] = mapped_column(DateTime)
    meeting_url: Mapped[Optional[str]] = mapped_column(String(500), nullable=True)
    max_participants: Mapped[Optional[int]] = mapped_column(Integer, nullable=True)
    created_by: Mapped[int] = mapped_column(ForeignKey("mitarbeiter.id"))


class TrainingRegistration(Base):
    __tablename__ = "mitarbeiter_schulungen"
    id: Mapped[int] = mapped_column(primary_key=True)
    training_id: Mapped[int] = mapped_column(ForeignKey("schulungen.id"))
    employee_id: Mapped[int] = mapped_column(ForeignKey("mitarbeiter.id"))
    registered_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)
    attended: Mapped[bool] = mapped_column(Boolean, default=False)


class SalesCoachMessage(Base):
    __tablename__ = "vertriebscoach_nachrichten"
    id: Mapped[int] = mapped_column(primary_key=True)
    employee_id: Mapped[int] = mapped_column(ForeignKey("mitarbeiter.id"))
    role: Mapped[str] = mapped_column(String(12))
    text: Mapped[str] = mapped_column(Text)
    created_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)


class TeamIn(BaseModel):
    name: str = Field(min_length=2, max_length=100)
    leader_id: Optional[int] = None

class MemberIn(BaseModel): team_id: int; employee_id: int
class ScheduleIn(BaseModel): employee_id: int; starts_at: datetime; ends_at: datetime; kind: str = "arbeit"; note: str = ""
class NewsIn(BaseModel): title: str; text: str; important: bool = False
class IncentiveIn(BaseModel): name: str; description: str = ""; minimum_contracts: int = Field(ge=0); reward_eur: float = Field(gt=0)
class ExpenseIn(BaseModel): category: str; description: str = ""; amount: float = Field(gt=0); spent_on: date = Field(default_factory=date.today)
class GoalIn(BaseModel): employee_id: Optional[int] = None; team_id: Optional[int] = None; period_start: date; period_end: date; target_contracts: int = Field(ge=0); target_revenue: float = Field(ge=0)
class ClosureIn(BaseModel):
    customer_id: Optional[int] = None
    customer_name: Optional[str] = Field(default=None, min_length=2, max_length=200)
    postal_code: Optional[str] = None
    phone: Optional[str] = None
    provider_id: Optional[int] = None
    tariff_id: Optional[int] = None
    contract_number: Optional[str] = Field(default=None, max_length=50)
    product: str = "strom"
    customer_kind: str = "privat"
    usage_kwh: float = Field(default=0, ge=0)
    completed_on: date = Field(default_factory=date.today)
    expected_commission: float = Field(default=0, ge=0)
    note: str = ""
    employee_id: Optional[int] = None
class ReviewIn(BaseModel): status: str; note: str = ""; provider_id: Optional[int] = None; tariff_id: Optional[int] = None; bracket_id: Optional[int] = None; usage_kwh: Optional[float] = Field(default=None, ge=0); expected_commission: Optional[float] = Field(default=None, ge=0)
class ProviderIn(BaseModel):
    name: str = Field(min_length=2, max_length=150)
    street: Optional[str] = None
    postal_code: Optional[str] = None
    city: Optional[str] = None
    phone: Optional[str] = None
    email: Optional[str] = None
    contact_person: Optional[str] = None
    notes: str = ""
    is_own_product: bool = False
class TariffIn(BaseModel):
    name: str = Field(min_length=1, max_length=200)
    external_id: Optional[str] = None
    product_type: Optional[Literal["strom", "gas"]] = None
    base_price_monthly: Optional[float] = Field(default=None, ge=0)
    price_per_kwh: Optional[float] = Field(default=None, ge=0)
    contract_term_months: Optional[int] = Field(default=None, ge=1)
    description: Optional[str] = None
class BracketIn(BaseModel):
    tier: int = Field(ge=1, le=3)
    usage_from: float = Field(ge=0, default=0)
    usage_to: Optional[float] = None
    commission_amount: float = Field(ge=0, default=0)
    commission_per_kwh: Optional[float] = None
class TariffImportIn(BaseModel):
    name: str
    external_id: Optional[str] = None
    brackets: list[BracketIn]
class ProviderImportIn(BaseModel):
    provider: str
    tariffs: list[TariffImportIn]
class DailyPerformanceIn(BaseModel): entry_date: date = Field(default_factory=date.today); contracts: int = Field(ge=0, le=100); cancellations: int = Field(ge=0, le=100); note: str = Field(default="", max_length=1000)
class TrainingIn(BaseModel): title: str = Field(min_length=3,max_length=200); description: str = ""; starts_at: datetime; ends_at: datetime; meeting_url: Optional[str] = None; max_participants: Optional[int] = Field(default=None,ge=1)
class PracticeIn(BaseModel): objection: str = Field(min_length=2,max_length=1000); product: str = "Stromvertrag"; customer_type: str = "Privatkunde"
class CoachChatIn(BaseModel): message: str = Field(min_length=2,max_length=2000)


def achievement(actual: float, target: float):
    percent = round(actual / target * 100, 1) if target else 0
    return {"actual": actual, "target": target, "percent": percent, "scale": "gruen" if percent >= 100 else "gelb" if percent >= 70 else "rot"}


@app.get("/api/agency/overview")
def overview(e: Employee = Depends(current), s: Session = Depends(db)):
    start = date.today().replace(day=1)
    closures = select(func.count(ClosureEntry.id)).where(ClosureEntry.status=="abgeschlossen", ClosureEntry.completed_on >= start)
    if e.role not in ("admin", "buchhaltung"): closures = closures.where(ClosureEntry.employee_id == e.id)
    commission = s.scalar(select(func.coalesce(func.sum(ClosureEntry.expected_commission), 0)).where(ClosureEntry.status=="abgeschlossen", ClosureEntry.completed_on >= start)) or 0
    expenses = s.scalar(select(func.coalesce(func.sum(Expense.amount), 0)).where(Expense.spent_on >= start)) or 0
    return {"month": start.isoformat(), "contracts": s.scalar(closures) or 0, "active_employees": s.scalar(select(func.count(Employee.id)).where(Employee.active.is_(True))) or 0, "commission_paid_out": commission, "expenses": expenses, "operating_result": round(commission-expenses, 2)}


@app.get("/api/agency/leaderboard")
def leaderboard(_: Employee = Depends(current), s: Session = Depends(db)):
    rows = s.execute(select(Employee.id, Employee.name, func.count(ClosureEntry.id).label("count")).outerjoin(ClosureEntry, (ClosureEntry.employee_id == Employee.id) & (ClosureEntry.status=="abgeschlossen")).where(Employee.active.is_(True)).group_by(Employee.id, Employee.name).order_by(func.count(ClosureEntry.id).desc())).all()
    return [{"employee_id": x.id, "name": x.name, "contracts": x.count} for x in rows]


@app.get("/api/agency/scorecards")
def scorecards(e: Employee = Depends(current), s: Session = Depends(db)):
    today = date.today(); first = today.replace(day=1)
    goals = list(s.scalars(select(SalesGoal).where(SalesGoal.period_start <= today, SalesGoal.period_end >= today)))
    employees = list(s.scalars(select(Employee).where(Employee.active.is_(True), Employee.role.not_in(("admin", "buchhaltung"))).order_by(Employee.name)))
    ids = visible_employee_ids(e, s)
    if ids is not None: employees = [x for x in employees if x.id in ids]
    result = []
    for employee in employees:
        contracts = s.scalar(select(func.count(ClosureEntry.id)).where(ClosureEntry.employee_id == employee.id, ClosureEntry.status == "abgeschlossen", ClosureEntry.completed_on >= first)) or 0
        commission = s.scalar(select(func.coalesce(func.sum(ClosureEntry.expected_commission), 0)).where(ClosureEntry.employee_id == employee.id, ClosureEntry.status == "abgeschlossen", ClosureEntry.completed_on >= first)) or 0
        goal = next((x for x in goals if x.employee_id == employee.id), None)
        result.append({"employee_id": employee.id, "name": employee.name, "tier": employee.tier, "contracts": achievement(contracts, goal.target_contracts if goal else 0), "revenue": achievement(commission, goal.target_revenue if goal else 0), "commission_rate": employee.commission_rate})
    return {"period_start": first.isoformat(), "employees": result}


@app.get("/api/agency/team-scorecards")
def team_scorecards(_: Employee = Depends(current), s: Session = Depends(db)):
    today = date.today(); first = today.replace(day=1); output=[]
    goals = list(s.scalars(select(SalesGoal).where(SalesGoal.period_start <= today, SalesGoal.period_end >= today, SalesGoal.team_id.is_not(None))))
    for team in s.scalars(select(Team).where(Team.active.is_(True))):
        members = [x.employee_id for x in s.scalars(select(TeamMember).where(TeamMember.team_id == team.id))]
        contracts = s.scalar(select(func.count(ClosureEntry.id)).where(ClosureEntry.employee_id.in_(members or [-1]), ClosureEntry.status == "abgeschlossen", ClosureEntry.completed_on >= first)) or 0
        goal = next((x for x in goals if x.team_id == team.id), None)
        output.append({"team_id": team.id, "name": team.name, "members": len(members), "contracts": achievement(contracts, goal.target_contracts if goal else 0)})
    return {"period_start": first.isoformat(), "teams": output}


@app.post("/api/goals")
def create_goal(data: GoalIn, e: Employee = Depends(admin), s: Session = Depends(db)):
    if not data.employee_id and not data.team_id: raise HTTPException(422, "Mitarbeiter oder Team erforderlich")
    if data.employee_id and data.team_id: raise HTTPException(422, "Nur Mitarbeiter oder Team wählen")
    if data.period_end < data.period_start: raise HTTPException(422, "Zeitraum ungültig")
    item=SalesGoal(**data.model_dump());s.add(item);log(s,e,"Vertriebsziel angelegt",str(item.employee_id or item.team_id));s.commit();return serialize(item)


@app.get("/api/employee/closures")
def my_closures(e: Employee = Depends(current), s: Session = Depends(db)):
    return [serialize(x) for x in s.scalars(select(ClosureEntry).where(ClosureEntry.employee_id == e.id).order_by(ClosureEntry.completed_on.desc(), ClosureEntry.id.desc()))]


@app.post("/api/employee/closures")
def submit_closure(data: ClosureIn, e: Employee = Depends(current), s: Session = Depends(db)):
    if e.role in ("admin", "buchhaltung"):
        if not data.employee_id: raise HTTPException(422, "employee_id erforderlich, wenn Admin für einen Mitarbeiter einträgt")
        owner_id = data.employee_id
    else:
        owner_id = e.id
    owner_employee = s.get(Employee, owner_id)
    customer_id = data.customer_id
    customer_name = data.customer_name
    if customer_id:
        cust = s.get(Customer, customer_id)
        if not cust: raise HTTPException(404, "Kunde nicht gefunden")
        if e.role not in ("admin", "buchhaltung") and cust.owner_id != e.id: raise HTTPException(403, "Keine Berechtigung")
        customer_name = cust.company or f"{cust.first_name or ''} {cust.last_name or ''}".strip()
    elif customer_name:
        parts = customer_name.split(" ", 1)
        cust = Customer(kind=data.customer_kind, first_name=parts[0], last_name=parts[1] if len(parts) > 1 else None, postal_code=data.postal_code or "", phone=data.phone, usage_kwh=data.usage_kwh, current_provider_id=data.provider_id, owner_id=owner_id)
        s.add(cust); s.flush()
        s.add(CustomerHistory(customer_id=cust.id, employee_id=e.id, detail="Kunde über Abschluss angelegt"))
        customer_id = cust.id
    else:
        raise HTTPException(422, "Kunde auswählen oder Kundenname angeben")
    expected_commission = data.expected_commission
    bracket_id = None
    if data.tariff_id:
        bracket = s.scalar(select(CommissionBracket).where(CommissionBracket.tariff_id == data.tariff_id, CommissionBracket.tier == owner_employee.tier, CommissionBracket.usage_from <= data.usage_kwh, (CommissionBracket.usage_to.is_(None)) | (CommissionBracket.usage_to >= data.usage_kwh)).order_by(CommissionBracket.usage_from.desc()))
        if bracket:
            bracket_id = bracket.id
            expected_commission = round(bracket.commission_amount + (bracket.commission_per_kwh or 0) * data.usage_kwh, 2)
    item = ClosureEntry(employee_id=owner_id, customer_id=customer_id, customer_name=customer_name, contract_number=data.contract_number, product=data.product, customer_kind=data.customer_kind, usage_kwh=data.usage_kwh, completed_on=data.completed_on, provider_id=data.provider_id, bracket_id=bracket_id, expected_commission=expected_commission, note=data.note)
    s.add(item); s.flush(); log(s, e, "Abschluss eingereicht", str(item.id))
    notify(s, "Neuer Abschluss zur Prüfung", f"{owner_employee.name}: {customer_name}", admins_only=True, link="provision")
    s.commit(); notify_update()
    return serialize(item)


def _parse_import_rows(filename: str, raw: bytes) -> list[dict]:
    ext = os.path.splitext(filename or "")[1].lower()
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        header = [str(h or "").strip().lower() for h in next(rows_iter)]
        out = []
        for row in rows_iter:
            if all(v is None for v in row): continue
            out.append({header[i]: row[i] for i in range(min(len(header), len(row)))})
        return out
    text = raw.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text), delimiter=";" if text.split("\n", 1)[0].count(";") >= text.split("\n", 1)[0].count(",") else ",")
    return [{(k or "").strip().lower(): v for k, v in row.items()} for row in reader]


@app.post("/api/admin/closures/import", dependencies=[Depends(require_module("csv_import"))])
def import_closures(file: UploadFile = File(...), e: Employee = Depends(admin), s: Session = Depends(db)):
    raw = file.file.read()
    try:
        rows = _parse_import_rows(file.filename or "", raw)
    except Exception as ex:
        raise HTTPException(422, f"Datei konnte nicht gelesen werden: {ex}")
    if not rows: raise HTTPException(422, "Keine Datenzeilen gefunden")

    imported, errors = 0, []
    for idx, row in enumerate(rows, start=2):  # row 1 = header
        try:
            emp_ref = str(row.get("mitarbeiter") or row.get("vp-nummer") or row.get("vp_nummer") or "").strip()
            if not emp_ref: raise ValueError("Mitarbeiter/VP-Nummer fehlt")
            owner = s.scalar(select(Employee).where(Employee.username == emp_ref))
            if not owner: owner = s.scalar(select(Employee).where(Employee.name.ilike(emp_ref)))
            if not owner: raise ValueError(f"Mitarbeiter '{emp_ref}' nicht gefunden")

            customer_name = str(row.get("kunde") or row.get("kundenname") or "").strip()
            if not customer_name: raise ValueError("Kundenname fehlt")

            product = str(row.get("produkt") or "strom").strip().lower()
            if product not in ("strom", "gas"): product = "strom"

            usage_raw = row.get("verbrauch") or row.get("verbrauch_kwh") or row.get("verbrauch (kwh)") or 0
            try: usage_kwh = float(str(usage_raw).replace(",", ".") or 0)
            except ValueError: usage_kwh = 0

            provider_id, tariff_id, bracket_id, expected_commission = None, None, None, 0.0
            provider_name = str(row.get("anbieter") or "").strip()
            if provider_name:
                provider = s.scalar(select(Provider).where(Provider.name.ilike(provider_name)))
                if provider:
                    provider_id = provider.id
                    tariff_name = str(row.get("tarif") or "").strip()
                    if tariff_name:
                        tariff = s.scalar(select(Tariff).where(Tariff.provider_id == provider.id, Tariff.name.ilike(tariff_name)))
                        if tariff:
                            tariff_id = tariff.id
                            bracket = s.scalar(select(CommissionBracket).where(CommissionBracket.tariff_id == tariff.id, CommissionBracket.tier == owner.tier, CommissionBracket.usage_from <= usage_kwh, (CommissionBracket.usage_to.is_(None)) | (CommissionBracket.usage_to >= usage_kwh)).order_by(CommissionBracket.usage_from.desc()))
                            if bracket:
                                bracket_id = bracket.id
                                expected_commission = round(bracket.commission_amount + (bracket.commission_per_kwh or 0) * usage_kwh, 2)

            date_raw = row.get("datum")
            completed_on = date.today()
            if date_raw:
                if isinstance(date_raw, datetime): completed_on = date_raw.date()
                elif isinstance(date_raw, date): completed_on = date_raw
                else:
                    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d.%m.%y"):
                        try: completed_on = datetime.strptime(str(date_raw).strip(), fmt).date(); break
                        except ValueError: continue

            item = ClosureEntry(
                employee_id=owner.id, customer_name=customer_name,
                contract_number=str(row.get("vertragsnummer") or "").strip() or None,
                product=product, usage_kwh=usage_kwh, completed_on=completed_on,
                provider_id=provider_id, bracket_id=bracket_id, expected_commission=expected_commission,
                note="Importiert per CSV/Excel",
            )
            s.add(item)
            imported += 1
        except Exception as ex:
            errors.append(f"Zeile {idx}: {ex}")

    s.commit()
    log(s, e, "Verträge importiert", f"{imported} erfolgreich, {len(errors)} Fehler")
    notify_update()
    return {"imported": imported, "errors": errors}


@app.get("/api/search")
def global_search(q: str, e: Employee = Depends(current), s: Session = Depends(db)):
    q = q.strip()
    if len(q) < 2: return {"customers": [], "employees": [], "closures": []}
    like = f"%{q}%"
    ids = visible_employee_ids(e, s)

    cust_stmt = select(Customer).where(or_(Customer.first_name.ilike(like), Customer.last_name.ilike(like), Customer.company.ilike(like), Customer.email.ilike(like), Customer.phone.ilike(like), Customer.postal_code.ilike(like)))
    if ids is not None: cust_stmt = cust_stmt.where(Customer.owner_id.in_(ids))
    customers = [{"id": c.id, "name": c.company or f"{c.first_name or ''} {c.last_name or ''}".strip(), "postal_code": c.postal_code, "status": c.status} for c in s.scalars(cust_stmt.limit(15))]

    employees = []
    if e.role in ("admin", "teamleiter", "buchhaltung"):
        emp_stmt = select(Employee).where(or_(Employee.name.ilike(like), Employee.username.ilike(like), Employee.email.ilike(like)))
        if ids is not None: emp_stmt = emp_stmt.where(Employee.id.in_(ids))
        employees = [{"id": x.id, "name": x.name, "username": x.username, "role": x.role} for x in s.scalars(emp_stmt.limit(15))]

    closure_stmt = select(ClosureEntry).where(or_(ClosureEntry.customer_name.ilike(like), ClosureEntry.contract_number.ilike(like)))
    if ids is not None: closure_stmt = closure_stmt.where(ClosureEntry.employee_id.in_(ids))
    closures = [{"id": x.id, "customer_name": x.customer_name, "contract_number": x.contract_number, "status": x.status, "completed_on": x.completed_on.isoformat()} for x in s.scalars(closure_stmt.order_by(ClosureEntry.completed_on.desc()).limit(15))]

    return {"customers": customers, "employees": employees, "closures": closures}


@app.get("/api/employee/performance-calendar")
def my_calendar(month: Optional[str] = None, e: Employee = Depends(current), s: Session = Depends(db)):
    start=date.fromisoformat(month+"-01") if month else date.today().replace(day=1)
    end=date(start.year + (start.month == 12), 1 if start.month == 12 else start.month+1, 1)
    rows=s.execute(select(ClosureEntry.completed_on, ClosureEntry.status, func.count(ClosureEntry.id), func.coalesce(func.sum(ClosureEntry.expected_commission),0)).where(ClosureEntry.employee_id==e.id,ClosureEntry.completed_on>=start,ClosureEntry.completed_on<end).group_by(ClosureEntry.completed_on,ClosureEntry.status).order_by(ClosureEntry.completed_on)).all()
    return [{"day":x[0].isoformat(),"status":x[1],"count":x[2],"expected_commission":x[3]} for x in rows]


@app.get("/api/admin/closures")
def all_closures(month: Optional[str] = None, _: Employee = Depends(admin), s: Session = Depends(db)):
    stmt=select(ClosureEntry)
    if month:
        start=date.fromisoformat(month+"-01"); end=date(start.year+(start.month==12),1 if start.month==12 else start.month+1,1); stmt=stmt.where(ClosureEntry.completed_on>=start,ClosureEntry.completed_on<end)
    return [serialize(x) for x in s.scalars(stmt.order_by(ClosureEntry.completed_on.desc(),ClosureEntry.id.desc()))]


STORNO_ALERT_ON = 40.0
STORNO_ALERT_OFF = 30.0

def commission_summary(rows):
    completed=[x for x in rows if x.status=="abgeschlossen"]; cancelled=[x for x in rows if x.status=="storno"]
    pending=[x for x in rows if x.status in ("eingereicht","bearbeitung","klaerung")]
    decided=len(completed)+len(cancelled)
    return {
        "total_commission": round(sum(x.expected_commission for x in completed),2),
        "pending_commission": round(sum(x.expected_commission for x in pending),2),
        "contracts_completed": len(completed), "contracts_pending": len(pending),
        "_cancelled": len(cancelled), "_rate": round(len(cancelled)/decided*100,1) if decided else 0.0,
    }

@app.get("/api/employee/commission-overview")
def my_commission_overview(e: Employee = Depends(current), s: Session = Depends(db)):
    rows=list(s.scalars(select(ClosureEntry).where(ClosureEntry.employee_id==e.id)))
    summary=commission_summary(rows)
    return {k:v for k,v in summary.items() if not k.startswith("_")}

@app.get("/api/admin/commission-overview")
def all_commission_overview(_: Employee = Depends(admin), s: Session = Depends(db)):
    result=[]
    for employee in s.scalars(select(Employee).where(Employee.active.is_(True)).order_by(Employee.name)):
        rows=list(s.scalars(select(ClosureEntry).where(ClosureEntry.employee_id==employee.id)))
        if not rows: continue
        summary=commission_summary(rows); rate=summary.pop("_rate"); cancelled=summary.pop("_cancelled")
        status=s.scalar(select(StornoStatus).where(StornoStatus.employee_id==employee.id))
        if not status: status=StornoStatus(employee_id=employee.id,alert=False); s.add(status)
        if not status.alert and rate>=STORNO_ALERT_ON: status.alert=True
        elif status.alert and rate<STORNO_ALERT_OFF: status.alert=False
        status.rate=rate; status.updated_at=datetime.utcnow()
        result.append({"employee_id":employee.id,"name":employee.name,"contracts_cancelled":cancelled,"cancellation_rate":rate,"storno_alert":status.alert, **summary})
    s.commit()
    return result

@app.get("/api/team-leaderboard")
def team_leaderboard(_: Employee = Depends(current), s: Session = Depends(db)):
    result=[]
    for employee in s.scalars(select(Employee).where(Employee.active.is_(True), Employee.role=="vertrieb").order_by(Employee.name)):
        rows=list(s.scalars(select(ClosureEntry).where(ClosureEntry.employee_id==employee.id)))
        summary=commission_summary(rows)
        result.append({"name":employee.name,"tier":employee.tier,"contracts_completed":summary["contracts_completed"],"contracts_cancelled":summary["_cancelled"],"cancellation_rate":summary["_rate"]})
    result.sort(key=lambda x:x["contracts_completed"],reverse=True)
    return result


@app.delete("/api/customers/{customer_id}")
def delete_customer(customer_id: int, e: Employee = Depends(current), s: Session = Depends(db)):
    c = s.get(Customer, customer_id)
    if not c: raise HTTPException(404, "Kunde nicht gefunden")
    if e.role != "admin" and c.owner_id != e.id: raise HTTPException(403, "Keine Berechtigung")
    for h in s.scalars(select(CustomerHistory).where(CustomerHistory.customer_id == customer_id)): s.delete(h)
    for t in s.scalars(select(Task).where(Task.customer_id == customer_id)): t.customer_id = None
    for cl in s.scalars(select(ClosureEntry).where(ClosureEntry.customer_id == customer_id)): cl.customer_id = None
    s.delete(c)
    log(s, e, "Kunde gelöscht", str(customer_id)); s.commit(); notify_update()
    return {"status": "deleted"}

@app.get("/api/employees/{employee_id}/kartei")
def employee_kartei(employee_id: int, e: Employee = Depends(current), s: Session = Depends(db)):
    if e.role != "admin" and e.id != employee_id: raise HTTPException(403, "Keine Berechtigung")
    target = s.get(Employee, employee_id)
    if not target: raise HTTPException(404, "Mitarbeiter nicht gefunden")
    closures = list(s.scalars(select(ClosureEntry).where(ClosureEntry.employee_id==employee_id).order_by(ClosureEntry.completed_on.desc())))
    summary = commission_summary(closures)
    teams = [s.get(Team, x.team_id).name for x in s.scalars(select(TeamMember).where(TeamMember.employee_id==employee_id)) if s.get(Team, x.team_id)]
    trainings = [serialize(s.get(Training, x.training_id)) for x in s.scalars(select(TrainingRegistration).where(TrainingRegistration.employee_id==employee_id)) if s.get(Training, x.training_id)]
    return {
        "employee": serialize_employee(target),
        "teams": teams,
        "customers": [serialize(x) for x in s.scalars(select(Customer).where(Customer.owner_id==employee_id).order_by(Customer.created_at.desc()))],
        "closures": [serialize(x) for x in closures],
        "commission_summary": {k: v for k, v in summary.items() if not k.startswith("_")},
        "tasks": [serialize(x) for x in s.scalars(select(Task).where(Task.assignee_id==employee_id).order_by(Task.due_date))],
        "daily_performance": [serialize(x) for x in s.scalars(select(DailyPerformance).where(DailyPerformance.employee_id==employee_id).order_by(DailyPerformance.entry_date.desc()))],
        "trainings": trainings,
    }


@app.post("/api/employees/{employee_id}/delete-account")
def delete_employee_account(employee_id: int, e: Employee = Depends(admin), s: Session = Depends(db)):
    """Account löschen: Login/TOTP/persönliche Daten werden entfernt, Kunden/Abschlüsse/Provisionshistorie bleiben für Buchhaltung erhalten."""
    x = s.get(Employee, employee_id)
    if not x: raise HTTPException(404, "Mitarbeiter nicht gefunden")
    if x.id == e.id: raise HTTPException(400, "Eigenen Account nicht löschen")
    x.active = False; x.name = "Ehemaliger Mitarbeiter"; x.email = None; x.phone = None; x.totp_secret = pyotp.random_base32()
    log(s, e, "Mitarbeiter-Account gelöscht (Geschäftsdaten bleiben erhalten)", x.username)
    s.commit()
    return {"status": "deleted", "mode": "account_only"}


@app.delete("/api/employees/{employee_id}")
def purge_employee(employee_id: int, e: Employee = Depends(admin), s: Session = Depends(db)):
    """Wirklich alles löschen: Mitarbeiter samt Kunden, Abschlüssen, Provisionen, Tagesmeldungen, Dokumenten, Terminen etc. unwiderruflich entfernen."""
    x = s.get(Employee, employee_id)
    if not x: raise HTTPException(404, "Mitarbeiter nicht gefunden")
    if x.id == e.id: raise HTTPException(400, "Eigenen Account nicht löschen")
    username = x.username
    customer_ids = list(s.scalars(select(Customer.id).where(Customer.owner_id == employee_id)))
    if customer_ids:
        for t in s.scalars(select(Task).where(Task.customer_id.in_(customer_ids))): t.customer_id = None
        for h in s.scalars(select(CustomerHistory).where(CustomerHistory.customer_id.in_(customer_ids))): s.delete(h)
    for row in s.scalars(select(Customer).where(Customer.owner_id == employee_id)): s.delete(row)
    for row in s.scalars(select(CustomerHistory).where(CustomerHistory.employee_id == employee_id)): s.delete(row)
    for row in s.scalars(select(Task).where(Task.assignee_id == employee_id)): s.delete(row)
    for row in s.scalars(select(Task).where(Task.creator_id == employee_id)): s.delete(row)
    for row in s.scalars(select(Activity).where(Activity.employee_id == employee_id)): s.delete(row)
    for row in s.scalars(select(ClosureEntry).where(ClosureEntry.employee_id == employee_id)): s.delete(row)
    for row in s.scalars(select(ClosureEntry).where(ClosureEntry.reviewed_by == employee_id)): row.reviewed_by = None
    for row in s.scalars(select(DailyPerformance).where(DailyPerformance.employee_id == employee_id)): s.delete(row)
    for row in s.scalars(select(TrainingRegistration).where(TrainingRegistration.employee_id == employee_id)): s.delete(row)
    for row in s.scalars(select(SalesCoachMessage).where(SalesCoachMessage.employee_id == employee_id)): s.delete(row)
    for row in s.scalars(select(TeamMember).where(TeamMember.employee_id == employee_id)): s.delete(row)
    for row in s.scalars(select(ScheduleEntry).where(ScheduleEntry.employee_id == employee_id)): s.delete(row)
    for row in s.scalars(select(StornoStatus).where(StornoStatus.employee_id == employee_id)): s.delete(row)
    for row in s.scalars(select(SalesGoal).where(SalesGoal.employee_id == employee_id)): s.delete(row)
    for row in s.scalars(select(Document).where(Document.owner_employee_id == employee_id)): s.delete(row)
    for row in s.scalars(select(Document).where(Document.uploaded_by == employee_id)): row.uploaded_by = e.id
    for row in s.scalars(select(Team).where(Team.leader_id == employee_id)): row.leader_id = None
    for row in s.scalars(select(News).where(News.author_id == employee_id)): row.author_id = e.id
    for row in s.scalars(select(Expense).where(Expense.created_by == employee_id)): row.created_by = e.id
    for row in s.scalars(select(Training).where(Training.created_by == employee_id)): row.created_by = e.id
    s.flush()
    s.delete(x)
    log(s, e, "Mitarbeiter vollständig gelöscht", username)
    s.commit()
    return {"status": "deleted", "mode": "full_purge"}


@app.get("/api/employees/{employee_id}/report.pdf")
def employee_report_pdf(employee_id: int, _: Employee = Depends(admin), s: Session = Depends(db)):
    target = s.get(Employee, employee_id)
    if not target: raise HTTPException(404, "Mitarbeiter nicht gefunden")
    closures = list(s.scalars(select(ClosureEntry).where(ClosureEntry.employee_id==employee_id).order_by(ClosureEntry.completed_on.desc())))
    summary = commission_summary(closures)
    status = s.scalar(select(StornoStatus).where(StornoStatus.employee_id==employee_id))
    lines = [
        f"VP-Nummer: {target.vp_nummer or '-'}   Stufe: {target.tier}   Rolle: {target.role}",
        "",
        f"Provision (abgeschlossen): {summary['total_commission']:.2f} EUR",
        f"Offene Provision: {summary['pending_commission']:.2f} EUR",
        f"Abgeschlossene Vertraege: {summary['contracts_completed']}",
        f"Offene/in Bearbeitung: {summary['contracts_pending']}",
        f"Stornoquote: {status.rate if status else 0:.1f} %",
        "",
        "Abschluesse:",
    ]
    for c in closures[:40]:
        lines.append(f"{c.completed_on}  {c.customer_name}  {c.status}  {c.expected_commission:.2f} EUR")
    name = f"report-{employee_id}-{uuid.uuid4().hex[:8]}.pdf"
    url = make_pdf(name, f"Mitarbeiterreport: {target.name}", lines)
    return Response((STORAGE / name).read_bytes(), media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="report-{target.username}.pdf"'})


@app.get("/api/employees/{employee_id}/report.xlsx")
def employee_report_xlsx(employee_id: int, _: Employee = Depends(admin), s: Session = Depends(db)):
    target = s.get(Employee, employee_id)
    if not target: raise HTTPException(404, "Mitarbeiter nicht gefunden")
    closures = list(s.scalars(select(ClosureEntry).where(ClosureEntry.employee_id==employee_id).order_by(ClosureEntry.completed_on.desc())))
    summary = commission_summary(closures)
    status = s.scalar(select(StornoStatus).where(StornoStatus.employee_id==employee_id))
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Übersicht"
    ws.append(["Mitarbeiter", target.name]); ws.append(["VP-Nummer", target.vp_nummer or "-"]); ws.append(["Stufe", target.tier])
    ws.append(["Provision (abgeschlossen)", summary["total_commission"]]); ws.append(["Offene Provision", summary["pending_commission"]])
    ws.append(["Abgeschlossene Verträge", summary["contracts_completed"]]); ws.append(["Offen/in Bearbeitung", summary["contracts_pending"]])
    ws.append(["Stornoquote %", status.rate if status else 0])
    ws2 = wb.create_sheet("Abschlüsse")
    ws2.append(["Datum", "Kunde", "Produkt", "Status", "Provision"])
    for c in closures:
        ws2.append([c.completed_on.isoformat(), c.customer_name, c.product, c.status, c.expected_commission])
    buf = io.BytesIO(); wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="report-{target.username}.xlsx"'})


@app.post("/api/admin/closures/{closure_id}/review")
def review_closure(closure_id: int, data: ReviewIn, e: Employee = Depends(admin), s: Session = Depends(db)):
    if data.status not in ("bearbeitung", "abgeschlossen", "storno", "klaerung"): raise HTTPException(422,"Ungültiger Prüfstatus")
    item=s.get(ClosureEntry,closure_id)
    if not item: raise HTTPException(404,"Abschluss nicht gefunden")
    old_vals={"status":item.status,"provider_id":item.provider_id,"expected_commission":item.expected_commission}
    item.status=data.status; item.note=(item.note+"\n"+data.note).strip(); item.reviewed_by=e.id; item.reviewed_at=datetime.utcnow()
    if data.provider_id: item.provider_id=data.provider_id
    if data.usage_kwh is not None: item.usage_kwh=data.usage_kwh
    bracket=None
    if data.bracket_id:
        bracket=s.get(CommissionBracket,data.bracket_id)
    elif data.tariff_id:
        owner=s.get(Employee,item.employee_id)
        bracket=s.scalar(select(CommissionBracket).where(CommissionBracket.tariff_id==data.tariff_id, CommissionBracket.tier==owner.tier, CommissionBracket.usage_from<=item.usage_kwh, (CommissionBracket.usage_to.is_(None))|(CommissionBracket.usage_to>=item.usage_kwh)).order_by(CommissionBracket.usage_from.desc()))
        if not bracket: raise HTTPException(422,f"Keine passende Provisionsstaffel für Tarif {data.tariff_id}, Stufe {owner.tier}, {item.usage_kwh} kWh gefunden")
    if bracket:
        item.bracket_id=bracket.id; item.expected_commission=round(bracket.commission_amount+(bracket.commission_per_kwh or 0)*item.usage_kwh,2)
    if data.expected_commission is not None: item.expected_commission=data.expected_commission
    new_vals={"status":item.status,"provider_id":item.provider_id,"expected_commission":item.expected_commission}
    log(s,e,"Abschluss geprüft",str(item.id),object_type="ClosureEntry",object_id=item.id,old=old_vals,new=new_vals)
    status_label={"abgeschlossen":"abgeschlossen (grün)","storno":"storniert","bearbeitung":"in Bearbeitung","klaerung":"Klärungsbedarf"}.get(item.status,item.status)
    notify(s,"Abschluss geprüft",f"{item.customer_name}: {status_label}",employee_id=item.employee_id,kind=("success" if item.status=="abgeschlossen" else "warning" if item.status=="storno" else "info"),link="dashboard")
    s.commit();notify_update();return serialize(item)


@app.delete("/api/admin/closures/{closure_id}")
def delete_closure(closure_id: int, e: Employee = Depends(admin), s: Session = Depends(db)):
    item=s.get(ClosureEntry,closure_id)
    if not item: raise HTTPException(404,"Abschluss nicht gefunden")
    s.delete(item); log(s,e,"Abschluss gelöscht",str(closure_id)); s.commit(); notify_update()
    return {"status":"deleted"}


@app.get("/api/providers")
def providers(_: Employee = Depends(current), s: Session = Depends(db)): return [serialize(x) for x in s.scalars(select(Provider).where(Provider.active.is_(True)).order_by(Provider.name))]
@app.post("/api/providers")
def create_provider(data: ProviderIn, e: Employee = Depends(admin), s: Session = Depends(db)):
    item=Provider(**data.model_dump()); s.add(item); s.flush(); log(s,e,"Anbieter angelegt",item.name); s.commit(); return serialize(item)
@app.get("/api/providers/{provider_id}/tariffs")
def provider_tariffs(provider_id: int, _: Employee = Depends(current), s: Session = Depends(db)): return [serialize(x) for x in s.scalars(select(Tariff).where(Tariff.provider_id==provider_id, Tariff.active.is_(True)).order_by(Tariff.name))]
@app.post("/api/providers/{provider_id}/tariffs")
def create_tariff(provider_id: int, data: TariffIn, e: Employee = Depends(admin), s: Session = Depends(db)):
    if not s.get(Provider,provider_id): raise HTTPException(404,"Anbieter nicht gefunden")
    item=Tariff(**data.model_dump(),provider_id=provider_id); s.add(item); log(s,e,"Tarif angelegt",f"{provider_id}:{item.name}"); s.commit(); return serialize(item)
@app.get("/api/tariffs/{tariff_id}/brackets")
def tariff_brackets(tariff_id: int, _: Employee = Depends(current), s: Session = Depends(db)): return [serialize(x) for x in s.scalars(select(CommissionBracket).where(CommissionBracket.tariff_id==tariff_id, CommissionBracket.active.is_(True)).order_by(CommissionBracket.tier,CommissionBracket.usage_from))]
@app.post("/api/tariffs/{tariff_id}/brackets")
def create_bracket(tariff_id: int, data: BracketIn, e: Employee = Depends(admin), s: Session = Depends(db)):
    if not s.get(Tariff,tariff_id): raise HTTPException(404,"Tarif nicht gefunden")
    item=CommissionBracket(**data.model_dump(),tariff_id=tariff_id); s.add(item); log(s,e,"Provisionsstaffel angelegt",str(tariff_id)); s.commit(); return serialize(item)


def get_or_create_own_provider(s: Session) -> Provider:
    p = s.scalar(select(Provider).where(Provider.is_own_product.is_(True)))
    if p: return p
    p = Provider(name="E1 Direktvertrieb", is_own_product=True, notes="Automatisch angelegt für eigene Produkte")
    s.add(p); s.flush()
    return p


class OwnProductIn(BaseModel):
    name: str = Field(min_length=1, max_length=200)
    product_type: Literal["strom", "gas"]
    base_price_monthly: float = Field(ge=0, default=0)
    price_per_kwh: float = Field(ge=0, default=0)
    contract_term_months: Optional[int] = Field(default=None, ge=1)
    description: Optional[str] = None


@app.get("/api/own-products")
def list_own_products(e: Employee = Depends(admin), s: Session = Depends(db)):
    provider = s.scalar(select(Provider).where(Provider.is_own_product.is_(True)))
    if not provider: return []
    return [serialize(x) for x in s.scalars(select(Tariff).where(Tariff.provider_id == provider.id).order_by(Tariff.name))]


@app.post("/api/own-products")
def create_own_product(data: OwnProductIn, e: Employee = Depends(admin), s: Session = Depends(db)):
    provider = get_or_create_own_provider(s)
    item = Tariff(provider_id=provider.id, **data.model_dump())
    s.add(item); log(s, e, "Eigenes Produkt angelegt", item.name); s.commit(); notify_update()
    return serialize(item)


@app.get("/api/public/own-products")
def public_own_products(s: Session = Depends(db)):
    if not module_enabled("eigene_produkte", s): return []
    rows = s.execute(select(Tariff, Provider.name).join(Provider, Provider.id == Tariff.provider_id).where(Provider.is_own_product.is_(True), Provider.active.is_(True), Tariff.active.is_(True)).order_by(Tariff.name)).all()
    return [{"id": t.id, "name": t.name, "provider_name": pname, "product_type": t.product_type, "base_price_monthly": t.base_price_monthly, "price_per_kwh": t.price_per_kwh, "contract_term_months": t.contract_term_months, "description": t.description} for t, pname in rows]


class OwnProductOrderIn(BaseModel):
    tariff_id: int
    name: str = Field(min_length=2, max_length=150)
    email: EmailStr
    phone: Optional[str] = None
    postal_code: str = Field(min_length=4, max_length=10)
    usage_kwh: float = Field(ge=0, default=0)


@app.post("/api/public/own-products/order")
def submit_own_product_order(data: OwnProductOrderIn, s: Session = Depends(db)):
    if not module_enabled("eigene_abschlussstrecke", s): raise HTTPException(403, "Diese Funktion ist aktuell nicht verfügbar.")
    tariff = s.get(Tariff, data.tariff_id)
    if not tariff or not tariff.active: raise HTTPException(404, "Tarif nicht gefunden")
    provider = s.get(Provider, tariff.provider_id)
    if not provider or not provider.is_own_product: raise HTTPException(422, "Kein eigenes Produkt")
    estimated = (tariff.base_price_monthly or 0) + (tariff.price_per_kwh or 0) * data.usage_kwh / 12
    item = OwnProductOrder(tariff_id=data.tariff_id, name=data.name, email=data.email, phone=data.phone, postal_code=data.postal_code, usage_kwh=data.usage_kwh, estimated_monthly_price=round(estimated, 2))
    s.add(item); s.commit(); notify_update()
    return {"status": "received", "estimated_monthly_price": item.estimated_monthly_price}


@app.get("/api/own-products/orders")
def list_own_product_orders(e: Employee = Depends(admin), s: Session = Depends(db)):
    rows = s.execute(select(OwnProductOrder, Tariff.name).join(Tariff, Tariff.id == OwnProductOrder.tariff_id).order_by(OwnProductOrder.created_at.desc())).all()
    return [{**serialize(o), "tariff_name": tname} for o, tname in rows]


@app.put("/api/own-products/orders/{order_id}")
def update_own_product_order(order_id: int, status: str = Form(...), e: Employee = Depends(admin), s: Session = Depends(db)):
    if status not in ("interessent", "bestaetigt", "aktiv", "storniert"): raise HTTPException(422, "Ungültiger Status")
    item = s.get(OwnProductOrder, order_id)
    if not item: raise HTTPException(404, "Bestellung nicht gefunden")
    item.status = status; log(s, e, "Eigene Bestellung aktualisiert", f"{order_id}:{status}")
    s.commit(); notify_update()
    return serialize(item)
@app.post("/api/providers/import")
def import_providers(data: list[ProviderImportIn], e: Employee = Depends(admin), s: Session = Depends(db)):
    providers_done=tariffs_done=brackets_done=0
    for p in data:
        provider=s.scalar(select(Provider).where(Provider.name==p.provider))
        if not provider:
            provider=Provider(name=p.provider); s.add(provider); s.flush(); providers_done+=1
        for t in p.tariffs:
            stmt=select(Tariff).where(Tariff.provider_id==provider.id, Tariff.external_id==t.external_id) if t.external_id else select(Tariff).where(Tariff.provider_id==provider.id, Tariff.name==t.name, Tariff.external_id.is_(None))
            tariff=s.scalar(stmt)
            if not tariff:
                tariff=Tariff(provider_id=provider.id,name=t.name,external_id=t.external_id); s.add(tariff); s.flush(); tariffs_done+=1
            else:
                for old in s.scalars(select(CommissionBracket).where(CommissionBracket.tariff_id==tariff.id)): s.delete(old)
                s.flush()
            for b in t.brackets:
                s.add(CommissionBracket(tariff_id=tariff.id,tier=b.tier,usage_from=b.usage_from,usage_to=b.usage_to,commission_amount=b.commission_amount,commission_per_kwh=b.commission_per_kwh)); brackets_done+=1
    log(s,e,"Anbieter-Import",f"{providers_done} Anbieter, {tariffs_done} Tarife, {brackets_done} Staffeln"); s.commit()
    return {"providers":providers_done,"tariffs":tariffs_done,"brackets":brackets_done}


@app.get("/api/admin/performance-calendar")
def admin_calendar(month: Optional[str] = None, _: Employee = Depends(admin), s: Session = Depends(db)):
    start=date.fromisoformat(month+"-01") if month else date.today().replace(day=1)
    end=date(start.year + (start.month == 12), 1 if start.month == 12 else start.month+1, 1)
    rows=s.execute(select(ClosureEntry.completed_on, Employee.id, Employee.name, ClosureEntry.status, func.count(ClosureEntry.id)).join(Employee,Employee.id==ClosureEntry.employee_id).where(ClosureEntry.completed_on>=start,ClosureEntry.completed_on<end).group_by(ClosureEntry.completed_on,Employee.id,Employee.name,ClosureEntry.status).order_by(ClosureEntry.completed_on,Employee.name)).all()
    return [{"day":x[0].isoformat(),"employee_id":x[1],"employee":x[2],"status":x[3],"count":x[4]} for x in rows]


@app.get("/api/employee/daily-performance")
def my_daily_performance(e: Employee = Depends(current), s: Session = Depends(db)):
    return [serialize(x) for x in s.scalars(select(DailyPerformance).where(DailyPerformance.employee_id == e.id).order_by(DailyPerformance.entry_date.desc()))]


@app.put("/api/employee/daily-performance")
def save_daily_performance(data: DailyPerformanceIn, e: Employee = Depends(current), s: Session = Depends(db)):
    if e.role in ("admin", "buchhaltung"): raise HTTPException(403, "Tagesmeldungen werden im Mitarbeiterportal erfasst")
    item=s.scalar(select(DailyPerformance).where(DailyPerformance.employee_id==e.id,DailyPerformance.entry_date==data.entry_date))
    if item:
        item.contracts=data.contracts;item.cancellations=data.cancellations;item.note=data.note;item.updated_at=datetime.utcnow()
    else:
        item=DailyPerformance(**data.model_dump(),employee_id=e.id);s.add(item)
    log(s,e,"Tagesmeldung gespeichert",data.entry_date.isoformat());s.commit();notify_update();return {**serialize(item),"net":item.contracts-item.cancellations}


@app.get("/api/admin/daily-performance")
def all_daily_performance(month: Optional[str] = None, _: Employee = Depends(admin), s: Session = Depends(db)):
    stmt=select(DailyPerformance,Employee).join(Employee,Employee.id==DailyPerformance.employee_id)
    if month:
        start=date.fromisoformat(month+"-01");end=date(start.year+(start.month==12),1 if start.month==12 else start.month+1,1);stmt=stmt.where(DailyPerformance.entry_date>=start,DailyPerformance.entry_date<end)
    rows=[]
    for entry, employee in s.execute(stmt.order_by(DailyPerformance.entry_date.desc())).all():
        data=serialize(entry);data.update({"employee":employee.name,"net":entry.contracts-entry.cancellations});rows.append(data)
    return rows


@app.get("/api/trainings")
def trainings(e: Employee = Depends(current), s: Session = Depends(db)):
    items=[]
    for training in s.scalars(select(Training).order_by(Training.starts_at)):
        data=serialize(training); data["registered"] = bool(s.scalar(select(TrainingRegistration.id).where(TrainingRegistration.training_id==training.id,TrainingRegistration.employee_id==e.id))); data["participants"] = s.scalar(select(func.count(TrainingRegistration.id)).where(TrainingRegistration.training_id==training.id)) or 0; items.append(data)
    return items


@app.post("/api/trainings")
def create_training(data: TrainingIn, e: Employee = Depends(admin), s: Session = Depends(db)):
    if data.ends_at <= data.starts_at: raise HTTPException(422,"Endzeit muss später sein")
    item=Training(**data.model_dump(),created_by=e.id);s.add(item);log(s,e,"Schulung angelegt",item.title);s.commit();return serialize(item)


@app.post("/api/trainings/{training_id}/register")
def register_training(training_id: int, employee_id: Optional[int] = None, e: Employee = Depends(current), s: Session = Depends(db)):
    target_id = employee_id if (employee_id and e.role=="admin") else e.id
    training=s.get(Training,training_id)
    if not training: raise HTTPException(404,"Schulung nicht gefunden")
    if s.scalar(select(TrainingRegistration.id).where(TrainingRegistration.training_id==training_id,TrainingRegistration.employee_id==target_id)): raise HTTPException(409,"Bereits angemeldet")
    count=s.scalar(select(func.count(TrainingRegistration.id)).where(TrainingRegistration.training_id==training_id)) or 0
    if training.max_participants and count>=training.max_participants: raise HTTPException(409,"Schulung ausgebucht")
    item=TrainingRegistration(training_id=training_id,employee_id=target_id);s.add(item);log(s,e,"Zu Schulung angemeldet",str(training_id));s.commit();return serialize(item)


PITCHES = [
    {"title":"Kurzer Strom-Pitch", "text":"Ich prüfe kostenlos, ob Ihr aktueller Stromtarif noch zu Ihrem Verbrauch passt. Wenn sich kein Vorteil ergibt, entsteht Ihnen keine Verpflichtung."},
    {"title":"Einwand: Kein Interesse", "text":"Verstehe ich. Darf ich nur eine kurze Frage stellen: Wann haben Sie Ihren Tarif zuletzt geprüft? Oft reicht ein kurzer Vergleich, damit Sie wissen, ob alles passt."},
    {"title":"Einwand: Ich habe schon einen Anbieter", "text":"Das ist gut. Gerade dann lohnt sich ein neutraler Vergleich vor Ablauf oder Preisänderung. Ich schaue nur, ob Ihr bestehender Tarif weiterhin sinnvoll ist."},
    {"title":"Einwand: Ich muss überlegen", "text":"Natürlich. Ich fasse die wichtigsten Punkte kurz zusammen und Sie entscheiden in Ruhe. Was wäre für Ihre Entscheidung noch offen?"},
]


@app.get("/api/training/pitches")
def pitches(_: Employee = Depends(current)): return PITCHES


def anthropic_client(api_key: str):
    from anthropic import Anthropic
    workspace_id = os.getenv("ANTHROPIC_WORKSPACE_ID")
    headers = {"anthropic-workspace-id": workspace_id} if workspace_id else None
    return Anthropic(api_key=api_key, default_headers=headers) if headers else Anthropic(api_key=api_key)

@app.post("/api/training/practice")
def practice(data: PracticeIn, _: Employee = Depends(current)):
    anthropic_key=os.getenv("ANTHROPIC_API_KEY")
    if anthropic_key and anthropic_key != "replace-with-a-new-rotated-key":
        try:
            result=anthropic_client(anthropic_key).messages.create(model=os.getenv("ANTHROPIC_MODEL","claude-sonnet-5"),max_tokens=500,system="Du bist ein deutschsprachiger, transparenter Vertriebscoach. Gib eine kurze Einwandbehandlung ohne Druck oder Preisversprechen.",messages=[{"role":"user","content":f"Produkt: {data.product}; Kunde: {data.customer_type}; Einwand: {data.objection}"}])
            return {"source":"claude","answer":result.content[0].text}
        except Exception as ex:
            print(f"[COACH ERROR] {type(ex).__name__}: {ex}", flush=True)
    match=next((x for x in PITCHES if "kein interesse" in data.objection.lower() and "Kein Interesse" in x["title"]),PITCHES[0])
    return {"source":"vorlage","answer":match["text"],"coach_tip":"Bleib freundlich, stelle nur eine offene Anschlussfrage und vermeide Druck."}


COACH_INSTRUCTIONS_BASE = "Du bist der EnergyOne Vertriebscoach von E1 Direktvertrieb. Hilf auf Deutsch bei Pitches, Einwandbehandlung, Gesprächsstruktur, Nachfass-Nachrichten, Selbstorganisation und Zielarbeit, und beantworte bei Bedarf auch allgemeine Fragen. Sei kurz, praktisch und respektvoll. Keine Druckmethoden, keine irreführenden Preisversprechen, keine Rechts- oder Steuerberatung. Frage bei fehlendem Kontext gezielt nach. Du hast Werkzeuge, um Aktionen direkt im System auszuführen (Kunden/Aufgaben anlegen, Abschluss einreichen; Admins zusätzlich News/Incentives/Ziele anlegen und E-Mails senden). Nutze sie nur, wenn der Nutzer erkennbar eine Aktion will, nicht bei reinen Fragen. Bevor du eine E-Mail tatsächlich versendest, lege Empfänger, Betreff und Text im Chat vor und warte auf eine ausdrückliche Bestätigung."

def coach_instructions_for(e: Employee) -> str:
    if e.role == "admin":
        persona = f"Du sprichst gerade mit {e.name}, Teamleitung/Admin bei E1 Direktvertrieb. Sprich sie/ihn mit Namen an, wenn es passt. Du darfst team- und agenturweite Themen besprechen (z.B. Teamzahlen, Mitarbeiterführung, Incentives), nicht nur Einzelverkauf."
    else:
        persona = f"Du sprichst gerade mit {e.name}, Vertriebsmitarbeiter/in (Stufe {e.tier}) bei E1 Direktvertrieb. Sprich sie/ihn mit Namen an, wenn es passt, und beziehe dich auf ihre/seine eigene Stufe/Provision, wenn relevant. Team- oder andere-Mitarbeiter-Daten bespricht du nicht mit ihr/ihm."
    return COACH_INSTRUCTIONS_BASE + "\n\n" + persona

def send_email(to: str, subject: str, body: str):
    host=os.getenv("SMTP_HOST")
    if not host: raise RuntimeError("SMTP ist nicht konfiguriert (SMTP_HOST fehlt).")
    msg=MIMEText(body,_charset="utf-8"); msg["Subject"]=subject; msg["From"]=os.getenv("SMTP_USER") or "no-reply@energyone.de"; msg["To"]=to
    with smtplib.SMTP(host, int(os.getenv("SMTP_PORT","587"))) as smtp:
        smtp.starttls()
        user=os.getenv("SMTP_USER")
        if user: smtp.login(user, os.getenv("SMTP_PASSWORD",""))
        smtp.send_message(msg)

@app.post("/api/public/apply")
def submit_application(name: str = Form(...), email: str = Form(...), phone: str = Form(""), message: str = Form(""), photo: Optional[UploadFile] = File(None), s: Session = Depends(db)):
    storage_name = None
    if photo is not None and photo.filename:
        ext = os.path.splitext(photo.filename)[1].lower() or ".jpg"
        if ext not in (".jpg", ".jpeg", ".png", ".webp"): raise HTTPException(422, "Foto muss JPG, PNG oder WEBP sein")
        raw = photo.file.read()
        if len(raw) > 8 * 1024 * 1024: raise HTTPException(422, "Foto darf maximal 8 MB groß sein")
        storage_name = f"apply-{uuid.uuid4().hex}{ext}"
        (STORAGE / storage_name).write_bytes(raw)
    item = JobApplication(name=name, email=email, phone=phone or None, message=message, photo_storage_name=storage_name)
    s.add(item); notify(s, "Neue Bewerbung", f"{name} hat sich beworben.", admins_only=True, link="mitarbeiter"); s.commit(); notify_update()
    recipients = os.getenv("APPLICATION_EMAIL", "luca.marrancone@gmail.com,saloorhan96@gmail.com")
    if recipients:
        body = f"Neue Bewerbung über die Website:\n\nName: {item.name}\nE-Mail: {item.email}\nTelefon: {item.phone or '-'}\n\nNachricht:\n{item.message or '-'}\n\nFoto im Portal unter Bewerbungen einsehbar."
        try: send_email(recipients, f"Neue Bewerbung: {item.name}", body)
        except Exception as ex: print(f"[APPLY EMAIL ERROR] {type(ex).__name__}: {ex}", flush=True)
    try:
        confirm_body = f"Hallo {item.name},\n\nvielen Dank für Ihre Bewerbung bei E1 Direktvertrieb. Wir haben sie erhalten und melden uns in Kürze persönlich bei Ihnen.\n\nViele Grüße\nIhr E1 Direktvertrieb Team"
        send_email(item.email, "Ihre Bewerbung bei E1 Direktvertrieb", confirm_body)
    except Exception as ex:
        print(f"[APPLY CONFIRM EMAIL ERROR] {type(ex).__name__}: {ex}", flush=True)
    return {"status": "received"}
@app.get("/api/admin/applications/{application_id}/photo")
def application_photo(application_id: int, e: Employee = Depends(admin), s: Session = Depends(db)):
    item = s.get(JobApplication, application_id)
    if not item or not item.photo_storage_name: raise HTTPException(404, "Kein Foto")
    p = STORAGE / item.photo_storage_name
    if not p.exists(): raise HTTPException(404, "Kein Foto")
    media = {".jpg":"image/jpeg",".jpeg":"image/jpeg",".png":"image/png",".webp":"image/webp"}.get(p.suffix.lower(), "application/octet-stream")
    return Response(p.read_bytes(), media_type=media)

class ApplicationReplyIn(BaseModel): message: str = Field(min_length=1)
@app.post("/api/admin/applications/{application_id}/reply")
def reply_application(application_id: int, data: ApplicationReplyIn, e: Employee = Depends(admin), s: Session = Depends(db)):
    item = s.get(JobApplication, application_id)
    if not item: raise HTTPException(404, "Bewerbung nicht gefunden")
    try:
        send_email(item.email, f"Antwort auf Ihre Bewerbung bei E1 Direktvertrieb", data.message)
    except Exception as ex:
        print(f"[APPLY REPLY ERROR] {type(ex).__name__}: {ex}", flush=True)
        raise HTTPException(500, f"E-Mail konnte nicht gesendet werden: {ex}")
    item.seen = True; log(s, e, "Bewerbung beantwortet", item.name); s.commit()
    return serialize(item)

AI_TOOLS_BASE = [
    {"name":"get_dashboard","description":"Zeigt die aktuellen Kennzahlen (Kunden, offene Aufgaben/Verträge/Rechnungen, Umsatz) für den eingeloggten Nutzer.","input_schema":{"type":"object","properties":{}}},
    {"name":"create_customer","description":"Legt einen neuen Kunden für den eingeloggten Mitarbeiter an.","input_schema":{"type":"object","properties":{"kind":{"type":"string","enum":["privat","firma"]},"first_name":{"type":"string"},"last_name":{"type":"string"},"company":{"type":"string"},"email":{"type":"string"},"phone":{"type":"string"},"postal_code":{"type":"string"},"street":{"type":"string"},"city":{"type":"string"},"usage_kwh":{"type":"number"}},"required":["kind","email","postal_code"]}},
    {"name":"create_task","description":"Legt eine Aufgabe für den eingeloggten Nutzer selbst an.","input_schema":{"type":"object","properties":{"title":{"type":"string"},"description":{"type":"string"},"due_date":{"type":"string","description":"ISO-Datum YYYY-MM-DD"}},"required":["title"]}},
    {"name":"submit_closure","description":"Reicht eine Abschluss-Meldung ein. Mitarbeiter reichen ihre eigene ein; Admins müssen employee_id angeben (z.B. zum abendlichen Nachtragen für einen Mitarbeiter).","input_schema":{"type":"object","properties":{"customer_name":{"type":"string"},"contract_number":{"type":"string"},"product":{"type":"string"},"customer_kind":{"type":"string"},"usage_kwh":{"type":"number"},"note":{"type":"string"},"employee_id":{"type":"integer","description":"Nur für Admins erforderlich"}},"required":["customer_name"]}},
    {"name":"classify_document","description":"Ordnet ein gerade hochgeladenes Dokument ein (Kategorie, ggf. Ablaufdatum/Betrag). Wird nach dem Ansehen eines im Chat mitgeschickten Dokuments aufgerufen.","input_schema":{"type":"object","properties":{"document_id":{"type":"integer"},"category":{"type":"string","enum":["ausweis","bankkarte","gewerbeanmeldung","fuehrungszeugnis","rechnung","sonstiges"]},"expires_on":{"type":"string","description":"ISO-Datum, falls erkennbar (z.B. Gültigkeit Führungszeugnis)"},"amount":{"type":"number","description":"Betrag, falls es eine Rechnung/Beleg ist"},"note":{"type":"string"}},"required":["document_id","category"]}},
]
AI_TOOLS_ADMIN = [
    {"name":"create_news","description":"Veröffentlicht eine News für alle Mitarbeiter.","input_schema":{"type":"object","properties":{"title":{"type":"string"},"text":{"type":"string"},"important":{"type":"boolean"}},"required":["title","text"]}},
    {"name":"create_incentive","description":"Legt ein neues Incentive an.","input_schema":{"type":"object","properties":{"name":{"type":"string"},"description":{"type":"string"},"minimum_contracts":{"type":"integer"},"reward_eur":{"type":"number"}},"required":["name","minimum_contracts","reward_eur"]}},
    {"name":"create_goal","description":"Setzt ein Vertriebsziel für einen Mitarbeiter oder ein Team.","input_schema":{"type":"object","properties":{"employee_id":{"type":"integer"},"team_id":{"type":"integer"},"period_start":{"type":"string"},"period_end":{"type":"string"},"target_contracts":{"type":"integer"},"target_revenue":{"type":"number"}},"required":["period_start","period_end"]}},
    {"name":"send_email","description":"Sendet eine E-Mail über den konfigurierten Mailserver. Vor dem tatsächlichen Versand IMMER zuerst im Chat Empfänger, Betreff und Text zur Bestätigung vorlegen und explizit auf Bestätigung warten.","input_schema":{"type":"object","properties":{"to":{"type":"string"},"subject":{"type":"string"},"body":{"type":"string"}},"required":["to","subject","body"]}},
    {"name":"create_employee","description":"Legt einen neuen Mitarbeiter mit Benutzername an; Antwort enthält den TOTP-Provisioning-Link zum Einrichten in Google Authenticator (im Admin-Bereich als QR anzeigbar).","input_schema":{"type":"object","properties":{"username":{"type":"string"},"name":{"type":"string"},"email":{"type":"string"},"role":{"type":"string","enum":["admin","vertrieb","support","buchhaltung"]}},"required":["username","name"]}},
    {"name":"reset_employee_totp","description":"Setzt das Google-Authenticator-TOTP eines Mitarbeiters zurück (z. B. bei Handyverlust) und liefert einen neuen Einrichtungslink.","input_schema":{"type":"object","properties":{"employee_id":{"type":"integer"}},"required":["employee_id"]}},
    {"name":"create_provider","description":"Legt einen neuen Energieanbieter (Stammdaten) an.","input_schema":{"type":"object","properties":{"name":{"type":"string"},"street":{"type":"string"},"postal_code":{"type":"string"},"city":{"type":"string"},"phone":{"type":"string"},"email":{"type":"string"},"contact_person":{"type":"string"},"notes":{"type":"string"}},"required":["name"]}},
    {"name":"create_tariff","description":"Legt einen neuen Tarif bei einem Anbieter an.","input_schema":{"type":"object","properties":{"provider_id":{"type":"integer"},"name":{"type":"string"},"external_id":{"type":"string"}},"required":["provider_id","name"]}},
    {"name":"create_bracket","description":"Legt eine Provisionsstaffel (Verbrauchsbereich + Provision) für einen Tarif und eine Stufe (1-3) an.","input_schema":{"type":"object","properties":{"tariff_id":{"type":"integer"},"tier":{"type":"integer"},"usage_from":{"type":"number"},"usage_to":{"type":"number"},"commission_amount":{"type":"number"},"commission_per_kwh":{"type":"number"}},"required":["tariff_id","tier","commission_amount"]}},
    {"name":"create_team","description":"Legt ein neues Team an.","input_schema":{"type":"object","properties":{"name":{"type":"string"},"leader_id":{"type":"integer"}},"required":["name"]}},
    {"name":"add_team_member","description":"Ordnet einen Mitarbeiter einem Team zu.","input_schema":{"type":"object","properties":{"team_id":{"type":"integer"},"employee_id":{"type":"integer"}},"required":["team_id","employee_id"]}},
    {"name":"create_schedule","description":"Legt einen Planungseintrag (Schicht/Termin) für einen Mitarbeiter an.","input_schema":{"type":"object","properties":{"employee_id":{"type":"integer"},"starts_at":{"type":"string","description":"ISO-Datetime"},"ends_at":{"type":"string","description":"ISO-Datetime"},"kind":{"type":"string"},"note":{"type":"string"}},"required":["employee_id","starts_at","ends_at"]}},
    {"name":"add_expense","description":"Erfasst eine Ausgabe.","input_schema":{"type":"object","properties":{"category":{"type":"string"},"description":{"type":"string"},"amount":{"type":"number"},"spent_on":{"type":"string","description":"ISO-Datum"}},"required":["category","amount"]}},
    {"name":"review_closure","description":"Prüft eine Mitarbeiter-Abschluss-Meldung: Status setzen (bearbeitung/abgeschlossen/storno/klaerung), optional Anbieter zuordnen. Bei tariff_id wird die Provisionsstaffel automatisch anhand der Stufe des Mitarbeiters und des Verbrauchs gewählt; bracket_id überschreibt das manuell.","input_schema":{"type":"object","properties":{"closure_id":{"type":"integer"},"status":{"type":"string","enum":["bearbeitung","abgeschlossen","storno","klaerung"]},"note":{"type":"string"},"provider_id":{"type":"integer"},"tariff_id":{"type":"integer"},"bracket_id":{"type":"integer"}},"required":["closure_id","status"]}},
    {"name":"create_training","description":"Legt eine Schulung an.","input_schema":{"type":"object","properties":{"title":{"type":"string"},"description":{"type":"string"},"starts_at":{"type":"string"},"ends_at":{"type":"string"},"meeting_url":{"type":"string"},"max_participants":{"type":"integer"}},"required":["title","starts_at","ends_at"]}},
    {"name":"rotate_master_key","description":"Generiert einen neuen Generalschlüssel (Notfallzugang) und gibt ihn im Chat aus. Wird NIE automatisch per E-Mail verschickt — der Admin gibt ihn selbst persönlich/telefonisch weiter.","input_schema":{"type":"object","properties":{}}},
]
AI_TOOLS_LOOKUP = [
    {"name":"search_tariffs","description":"Sucht Anbieter/Tarife nach Namen und liefert deren Provisionsstaffeln über alle Stufen (1-3). Für Fragen wie 'Was ist die Provision bei <Anbieter> <Tarif> in Stufe 2 bei X kWh?'.","input_schema":{"type":"object","properties":{"query":{"type":"string","description":"Anbieter- oder Tarifname (Teilstring)"}},"required":["query"]}},
]
def ai_tools_for(role: str): return AI_TOOLS_BASE + AI_TOOLS_LOOKUP + (AI_TOOLS_ADMIN if role=="admin" else [])

def run_ai_tool(tool_name: str, tool_input: dict, e: Employee, s: Session):
    try:
        if tool_name=="get_dashboard":
            from .app import dashboard
            return dashboard(e=e, s=s)
        if tool_name=="create_customer":
            data=CustomerIn(kind=tool_input.get("kind","privat"), first_name=tool_input.get("first_name"), last_name=tool_input.get("last_name"), company=tool_input.get("company"), email=tool_input.get("email"), phone=tool_input.get("phone"), postal_code=tool_input.get("postal_code",""), street=tool_input.get("street"), city=tool_input.get("city"), usage_kwh=tool_input.get("usage_kwh",0))
            result=create_customer(data,e,s); log(s,e,"KI: Kunde angelegt",str(result.get("id"))); return result
        if tool_name=="create_task":
            data=TaskIn(title=tool_input["title"], description=tool_input.get("description",""), assignee_id=e.id, due_date=date.fromisoformat(tool_input["due_date"]) if tool_input.get("due_date") else None)
            result=create_task(data,e,s); log(s,e,"KI: Aufgabe angelegt",data.title); return result
        if tool_name=="submit_closure":
            data=ClosureIn(customer_name=tool_input["customer_name"], contract_number=tool_input.get("contract_number"), product=tool_input.get("product","strom"), customer_kind=tool_input.get("customer_kind","privat"), usage_kwh=tool_input.get("usage_kwh",0), note=tool_input.get("note",""), employee_id=tool_input.get("employee_id"))
            result=submit_closure(data,e,s); log(s,e,"KI: Abschluss eingereicht",str(result.get("id"))); return result
        if tool_name=="search_tariffs":
            q=f"%{tool_input['query']}%"
            rows=list(s.execute(select(Tariff,Provider.name).join(Provider,Provider.id==Tariff.provider_id).where((Provider.name.ilike(q))|(Tariff.name.ilike(q))).limit(10)).all())
            out=[]
            for t,provider_name in rows:
                brackets=list(s.scalars(select(CommissionBracket).where(CommissionBracket.tariff_id==t.id).order_by(CommissionBracket.tier,CommissionBracket.usage_from)))
                out.append({"provider":provider_name,"tariff":t.name,"external_id":t.external_id,"brackets":[{"stufe":b.tier,"von_kwh":b.usage_from,"bis_kwh":b.usage_to,"provision_eur":b.commission_amount,"provision_je_kwh":b.commission_per_kwh} for b in brackets]})
            return out or {"info":"Keine Treffer."}
        if tool_name=="classify_document":
            doc=s.get(Document,tool_input["document_id"])
            if not doc: return {"error":"Dokument nicht gefunden"}
            if e.role!="admin" and doc.owner_employee_id!=e.id: return {"error":"Keine Berechtigung für dieses Dokument"}
            doc.category=tool_input["category"]
            if tool_input.get("expires_on"): doc.expires_on=date.fromisoformat(tool_input["expires_on"])
            if tool_input.get("amount") is not None: doc.amount=tool_input["amount"]
            if tool_input.get("note"): doc.note=tool_input["note"]
            log(s,e,"KI: Dokument eingeordnet",f"{doc.category}:{doc.id}"); s.commit(); return serialize(doc)
        if e.role!="admin":
            return {"error":"Diese Aktion ist nur für Admins verfügbar."}
        if tool_name=="create_news":
            data=NewsIn(title=tool_input["title"], text=tool_input["text"], important=tool_input.get("important",False))
            result=publish_news(data,e,s); log(s,e,"KI: News veröffentlicht",data.title); return result
        if tool_name=="create_incentive":
            data=IncentiveIn(name=tool_input["name"], description=tool_input.get("description",""), minimum_contracts=tool_input.get("minimum_contracts",0), reward_eur=tool_input["reward_eur"])
            result=create_incentive(data,e,s); log(s,e,"KI: Incentive angelegt",data.name); return result
        if tool_name=="create_goal":
            data=GoalIn(employee_id=tool_input.get("employee_id"), team_id=tool_input.get("team_id"), period_start=date.fromisoformat(tool_input["period_start"]), period_end=date.fromisoformat(tool_input["period_end"]), target_contracts=tool_input.get("target_contracts",0), target_revenue=tool_input.get("target_revenue",0))
            result=create_goal(data,e,s); log(s,e,"KI: Ziel angelegt",str(result.get("id"))); return result
        if tool_name=="send_email":
            send_email(tool_input["to"], tool_input["subject"], tool_input["body"]); log(s,e,"KI: E-Mail gesendet",tool_input["to"]); return {"status":"gesendet","to":tool_input["to"]}
        if tool_name=="create_employee":
            data=EmployeeIn(username=tool_input["username"], name=tool_input["name"], email=tool_input.get("email"), role=tool_input.get("role","vertrieb"))
            result=create_employee(data,e,s); log(s,e,"KI: Mitarbeiter angelegt",data.username); return result
        if tool_name=="reset_employee_totp":
            result=reset_totp(tool_input["employee_id"],e,s); log(s,e,"KI: TOTP zurückgesetzt",str(tool_input["employee_id"])); return result
        if tool_name=="create_provider":
            data=ProviderIn(name=tool_input["name"], street=tool_input.get("street"), postal_code=tool_input.get("postal_code"), city=tool_input.get("city"), phone=tool_input.get("phone"), email=tool_input.get("email"), contact_person=tool_input.get("contact_person"), notes=tool_input.get("notes",""))
            result=create_provider(data,e,s); log(s,e,"KI: Anbieter angelegt",data.name); return result
        if tool_name=="create_tariff":
            data=TariffIn(name=tool_input["name"], external_id=tool_input.get("external_id"))
            result=create_tariff(tool_input["provider_id"],data,e,s); log(s,e,"KI: Tarif angelegt",data.name); return result
        if tool_name=="create_bracket":
            data=BracketIn(tier=tool_input["tier"], usage_from=tool_input.get("usage_from",0), usage_to=tool_input.get("usage_to"), commission_amount=tool_input.get("commission_amount",0), commission_per_kwh=tool_input.get("commission_per_kwh"))
            result=create_bracket(tool_input["tariff_id"],data,e,s); log(s,e,"KI: Staffel angelegt",str(tool_input["tariff_id"])); return result
        if tool_name=="create_team":
            data=TeamIn(name=tool_input["name"], leader_id=tool_input.get("leader_id"))
            result=create_team(data,e,s); log(s,e,"KI: Team angelegt",data.name); return result
        if tool_name=="add_team_member":
            data=MemberIn(team_id=tool_input["team_id"], employee_id=tool_input["employee_id"])
            result=add_member(data,e,s); log(s,e,"KI: Teammitglied zugeordnet",str(data.employee_id)); return result
        if tool_name=="create_schedule":
            data=ScheduleIn(employee_id=tool_input["employee_id"], starts_at=datetime.fromisoformat(tool_input["starts_at"]), ends_at=datetime.fromisoformat(tool_input["ends_at"]), kind=tool_input.get("kind","arbeit"), note=tool_input.get("note",""))
            result=schedule(data,e,s); log(s,e,"KI: Planung erstellt",str(data.employee_id)); return result
        if tool_name=="add_expense":
            data=ExpenseIn(category=tool_input["category"], description=tool_input.get("description",""), amount=tool_input["amount"], spent_on=date.fromisoformat(tool_input["spent_on"]) if tool_input.get("spent_on") else date.today())
            result=add_expense(data,e,s); log(s,e,"KI: Ausgabe erfasst",data.category); return result
        if tool_name=="review_closure":
            data=ReviewIn(status=tool_input["status"], note=tool_input.get("note",""), provider_id=tool_input.get("provider_id"), tariff_id=tool_input.get("tariff_id"), bracket_id=tool_input.get("bracket_id"))
            result=review_closure(tool_input["closure_id"],data,e,s); log(s,e,"KI: Abschluss geprüft",str(tool_input["closure_id"])); return result
        if tool_name=="create_training":
            data=TrainingIn(title=tool_input["title"], description=tool_input.get("description",""), starts_at=datetime.fromisoformat(tool_input["starts_at"]), ends_at=datetime.fromisoformat(tool_input["ends_at"]), meeting_url=tool_input.get("meeting_url"), max_participants=tool_input.get("max_participants"))
            result=create_training(data,e,s); log(s,e,"KI: Schulung angelegt",data.title); return result
        if tool_name=="rotate_master_key":
            result=rotate_master_key(MasterKeyIn(new_key=None),e,s); log(s,e,"KI: Generalschlüssel geändert"); return result
        return {"error":f"Unbekanntes Tool: {tool_name}"}
    except HTTPException as ex:
        return {"error": ex.detail}
    except Exception as ex:
        return {"error": str(ex)}

@app.get("/api/training/coach/history")
def coach_history(e: Employee = Depends(current), s: Session = Depends(db)):
    return [serialize(x) for x in s.scalars(select(SalesCoachMessage).where(SalesCoachMessage.employee_id==e.id).order_by(SalesCoachMessage.created_at.desc()).limit(30))][::-1]

@app.delete("/api/training/coach/history")
def clear_coach_history(e: Employee = Depends(current), s: Session = Depends(db)):
    for m in s.scalars(select(SalesCoachMessage).where(SalesCoachMessage.employee_id==e.id)): s.delete(m)
    s.commit()
    return {"status": "cleared"}

COACH_FALLBACK_TEXTS = {
    "Der KI-Coach ist gerade nicht erreichbar. Nutze bis dahin die Pitch-Vorlagen oder frage deine Teamleitung.",
    "Für individuelle Antworten bitte die KI-Anbindung aktivieren. Bis dahin: Beschreibe den Einwand kurz, bleib freundlich und stelle eine offene Frage.",
    "Die Anfrage war zu umfangreich für eine direkte Antwort, bitte präzisiere sie.",
}

def run_coach_loop(user_content, e: Employee, s: Session):
    history_raw=list(s.scalars(select(SalesCoachMessage).where(SalesCoachMessage.employee_id==e.id).order_by(SalesCoachMessage.created_at.desc()).limit(20)))[::-1]
    history=[x for x in history_raw[:-1] if x.text not in COACH_FALLBACK_TEXTS][-12:]
    anthropic_key=os.getenv("ANTHROPIC_API_KEY")
    print(f"[COACH DEBUG] anthropic_key_set={bool(anthropic_key)} openai_key_set={bool(os.getenv('OPENAI_API_KEY'))} workspace_id_set={bool(os.getenv('ANTHROPIC_WORKSPACE_ID'))}", flush=True)
    if anthropic_key and anthropic_key != "replace-with-a-new-rotated-key":
        try:
            client=anthropic_client(anthropic_key)
            messages=[{"role":x.role,"content":x.text} for x in history] + [{"role":"user","content":user_content}]
            tools=ai_tools_for(e.role)
            answer=None
            for _ in range(5):
                response=client.messages.create(model=os.getenv("ANTHROPIC_MODEL","claude-sonnet-5"),max_tokens=800,system=coach_instructions_for(e),messages=messages,tools=tools)
                if response.stop_reason!="tool_use":
                    answer="".join(b.text for b in response.content if b.type=="text")
                    break
                messages.append({"role":"assistant","content":response.content})
                tool_results=[]
                for block in response.content:
                    if block.type=="tool_use":
                        result=run_ai_tool(block.name, block.input, e, s)
                        tool_results.append({"type":"tool_result","tool_use_id":block.id,"content":json.dumps(result, default=str)})
                messages.append({"role":"user","content":tool_results})
            if answer is None: answer="Die Anfrage war zu umfangreich für eine direkte Antwort, bitte präzisiere sie."
        except Exception as ex:
            import traceback
            print(f"[COACH ERROR] {type(ex).__name__}: {ex}", flush=True)
            traceback.print_exc()
            answer="Der KI-Coach ist gerade nicht erreichbar. Nutze bis dahin die Pitch-Vorlagen oder frage deine Teamleitung."
    else:
        answer="Für individuelle Antworten bitte die KI-Anbindung aktivieren. Bis dahin: Beschreibe den Einwand kurz, bleib freundlich und stelle eine offene Frage."
    reply=SalesCoachMessage(employee_id=e.id,role="assistant",text=answer);s.add(reply);s.commit();return serialize(reply)

@app.post("/api/training/coach/chat")
def coach_chat(data: CoachChatIn, e: Employee = Depends(current), s: Session = Depends(db)):
    s.add(SalesCoachMessage(employee_id=e.id,role="user",text=data.message));s.flush()
    return run_coach_loop(data.message, e, s)

DOCUMENT_MEDIA_TYPES = {".pdf":"application/pdf", ".png":"image/png", ".jpg":"image/jpeg", ".jpeg":"image/jpeg", ".webp":"image/webp"}

@app.post("/api/training/coach/upload")
def coach_upload(message: str = Form(""), owner_employee_id: Optional[int] = Form(None), file: UploadFile = File(...), e: Employee = Depends(current), s: Session = Depends(db)):
    ext = os.path.splitext(file.filename or "")[1].lower()
    raw = file.file.read()
    owner = owner_employee_id if (e.role == "admin" and owner_employee_id) else (None if e.role == "admin" else e.id)
    storage_name = f"doc-{uuid.uuid4().hex}-{file.filename}"
    (STORAGE / storage_name).write_bytes(raw)
    doc = Document(owner_employee_id=owner, category="unsortiert", filename=file.filename, storage_name=storage_name, uploaded_by=e.id)
    s.add(doc); s.flush(); doc_id = doc.id
    text_note = f"[Datei hochgeladen: {file.filename}, document_id={doc_id}] {message}".strip()
    s.add(SalesCoachMessage(employee_id=e.id, role="user", text=text_note)); s.flush()

    media_type = DOCUMENT_MEDIA_TYPES.get(ext)
    if media_type:
        block_type = "document" if media_type == "application/pdf" else "image"
        content = [
            {"type": block_type, "source": {"type": "base64", "media_type": media_type, "data": base64.b64encode(raw).decode()}},
            {"type": "text", "text": f"{text_note}\n\nBitte sieh dir das Dokument an und ordne es über das classify_document-Tool ein (document_id={doc_id})."},
        ]
    else:
        content = f"{text_note}\n\n(Dateityp {ext or 'unbekannt'} kann nicht direkt angesehen werden — bitte über classify_document mit document_id={doc_id} manuell einordnen, falls erkennbar.)"
    s.commit()
    return run_coach_loop(content, e, s)


@app.get("/api/teams")
def teams(_: Employee = Depends(current), s: Session = Depends(db)): return [serialize(x) for x in s.scalars(select(Team).order_by(Team.name))]
@app.post("/api/teams")
def create_team(data: TeamIn, e: Employee = Depends(admin), s: Session = Depends(db)):
    item=Team(**data.model_dump()); s.add(item); s.flush(); log(s,e,"Team angelegt",item.name); s.commit(); return serialize(item)
@app.post("/api/teams/members")
def add_member(data: MemberIn, e: Employee = Depends(admin), s: Session = Depends(db)):
    if not s.get(Team,data.team_id) or not s.get(Employee,data.employee_id): raise HTTPException(404,"Team oder Mitarbeiter nicht gefunden")
    item=TeamMember(**data.model_dump());s.add(item);log(s,e,"Teammitglied zugeordnet",str(data.employee_id));s.commit();return serialize(item)


@app.get("/api/planning")
def planning(e: Employee = Depends(current), s: Session = Depends(db)):
    stmt=select(ScheduleEntry) if e.role=="admin" else select(ScheduleEntry).where(ScheduleEntry.employee_id==e.id)
    return [serialize(x) for x in s.scalars(stmt.order_by(ScheduleEntry.starts_at))]
@app.post("/api/planning")
def schedule(data: ScheduleIn, e: Employee = Depends(current), s: Session = Depends(db)):
    if e.role!="admin" and data.employee_id!=e.id: raise HTTPException(403,"Keine Berechtigung")
    if data.ends_at<=data.starts_at: raise HTTPException(422,"Endzeit muss später sein")
    item=ScheduleEntry(**data.model_dump());s.add(item);log(s,e,"Planung erstellt",str(data.employee_id));s.commit();return serialize(item)


@app.get("/api/news")
def news(_: Employee = Depends(current), s: Session = Depends(db)): return [serialize(x) for x in s.scalars(select(News).order_by(News.important.desc(),News.published_at.desc()))]
@app.post("/api/news")
def publish_news(data: NewsIn, e: Employee = Depends(admin), s: Session = Depends(db)):
    item=News(**data.model_dump(),author_id=e.id);s.add(item);log(s,e,"News veröffentlicht",item.title);s.commit();notify_update();return serialize(item)
@app.delete("/api/news/{news_id}")
def delete_news(news_id: int, e: Employee = Depends(admin), s: Session = Depends(db)):
    item=s.get(News,news_id)
    if not item: raise HTTPException(404,"News nicht gefunden")
    s.delete(item); log(s,e,"News gelöscht",item.title); s.commit(); notify_update()
    return {"status":"deleted"}
_ENERGY_NEWS_CACHE = {"items": [], "fetched_at": 0.0}
_ENERGY_NEWS_TTL = 1800
_ENERGY_NEWS_FEED = "https://news.google.com/rss/search?q=Strom%20OR%20Gas%20OR%20Energiepreise%20Deutschland%20when:14d&hl=de&gl=DE&ceid=DE:de"

def _fetch_energy_news():
    import time, urllib.request, xml.etree.ElementTree as ET
    now = time.time()
    if now - _ENERGY_NEWS_CACHE["fetched_at"] < _ENERGY_NEWS_TTL and _ENERGY_NEWS_CACHE["items"]:
        return _ENERGY_NEWS_CACHE["items"]
    try:
        req = urllib.request.Request(_ENERGY_NEWS_FEED, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=8) as resp:
            xml_data = resp.read()
        root = ET.fromstring(xml_data)
        items = []
        for item in root.findall(".//item")[:10]:
            title = (item.findtext("title") or "").strip()
            link = (item.findtext("link") or "").strip()
            pub_date = (item.findtext("pubDate") or "").strip()
            source_el = item.find("source")
            source = source_el.text.strip() if source_el is not None and source_el.text else ""
            if title and link:
                items.append({"title": title, "link": link, "source": source, "published": pub_date})
        if items:
            _ENERGY_NEWS_CACHE["items"] = items
            _ENERGY_NEWS_CACHE["fetched_at"] = now
        return _ENERGY_NEWS_CACHE["items"]
    except Exception as ex:
        print(f"[ENERGY NEWS ERROR] {type(ex).__name__}: {ex}", flush=True)
        return _ENERGY_NEWS_CACHE["items"]

@app.get("/api/public/energy-news")
def public_energy_news():
    return _fetch_energy_news()


@app.get("/api/incentives")
def incentives(_: Employee = Depends(current), s: Session = Depends(db)): return [serialize(x) for x in s.scalars(select(Incentive).order_by(Incentive.id.desc()))]
@app.post("/api/incentives")
def create_incentive(data: IncentiveIn, e: Employee = Depends(admin), s: Session = Depends(db)):
    item=Incentive(**data.model_dump());s.add(item);log(s,e,"Incentive angelegt",item.name);s.commit();return serialize(item)


@app.get("/api/expenses")
def expenses(_: Employee = Depends(admin), s: Session = Depends(db)): return [serialize(x) for x in s.scalars(select(Expense).order_by(Expense.spent_on.desc()))]
@app.post("/api/expenses")
def add_expense(data: ExpenseIn, e: Employee = Depends(admin), s: Session = Depends(db)):
    item=Expense(**data.model_dump(),created_by=e.id);s.add(item);log(s,e,"Ausgabe erfasst",item.category);s.commit();return serialize(item)


@app.get("/api/documents")
def documents(owner_employee_id: Optional[int] = None, category: Optional[str] = None, e: Employee = Depends(current), s: Session = Depends(db)):
    stmt = select(Document)
    if e.role != "admin":
        stmt = stmt.where(Document.owner_employee_id == e.id)
    elif owner_employee_id is not None:
        stmt = stmt.where(Document.owner_employee_id == owner_employee_id)
    if category: stmt = stmt.where(Document.category == category)
    return [serialize(x) for x in s.scalars(stmt.order_by(Document.uploaded_at.desc()))]


@app.post("/api/documents")
def upload_document(category: str = Form(...), owner_employee_id: Optional[int] = Form(None), amount: Optional[float] = Form(None), paid: Optional[bool] = Form(None), expires_on: Optional[str] = Form(None), note: str = Form(""), file: UploadFile = File(...), e: Employee = Depends(current), s: Session = Depends(db)):
    owner = owner_employee_id if (e.role == "admin" and owner_employee_id) else (None if e.role == "admin" else e.id)
    storage_name = f"doc-{uuid.uuid4().hex}-{file.filename}"
    (STORAGE / storage_name).write_bytes(file.file.read())
    item = Document(owner_employee_id=owner, category=category, filename=file.filename, storage_name=storage_name, amount=amount, paid=paid, expires_on=date.fromisoformat(expires_on) if expires_on else None, note=note, uploaded_by=e.id)
    s.add(item); s.flush(); log(s, e, "Dokument hochgeladen", f"{category}:{item.id}"); s.commit()
    return serialize(item)


@app.get("/api/documents/{document_id}/file")
def download_document(document_id: int, e: Employee = Depends(current), s: Session = Depends(db)):
    item = s.get(Document, document_id)
    if not item: raise HTTPException(404, "Dokument nicht gefunden")
    if e.role != "admin" and item.owner_employee_id != e.id: raise HTTPException(403, "Keine Berechtigung")
    path = STORAGE / item.storage_name
    if not path.exists(): raise HTTPException(404, "Datei nicht gefunden")
    return Response(path.read_bytes(), media_type="application/octet-stream", headers={"Content-Disposition": f'attachment; filename="{item.filename}"'})


@app.get("/api/admin/documents/expiring")
def documents_expiring(days: int = 60, _: Employee = Depends(admin), s: Session = Depends(db)):
    cutoff = date.today() + timedelta(days=days)
    rows = s.execute(select(Document, Employee.name).outerjoin(Employee, Employee.id == Document.owner_employee_id).where(Document.expires_on.is_not(None), Document.expires_on <= cutoff).order_by(Document.expires_on)).all()
    return [{**serialize(d), "employee_name": name} for d, name in rows]
